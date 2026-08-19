# -*- coding: utf-8 -*-
"""通用 Markdown → Word 转换器（复用 md_to_docx.py 的风格体系）。

用法：
    python scripts/md_to_docx_general.py <input.md> <output.docx> --title 标题 --version 版本 [--date 日期]

特性：
- 封面 + Word 自动目录（TOC 域，Word 打开时更新）
- 标题/段落/列表/引用/表格/代码块/行内 **粗体** 与 `代码`
- ASCII 框线图（┌─┐ 网格，如层次架构图、帧结构图）→ Word 原生表格（带底纹）
- ASCII 时序图（│ 生命线 + →/← 箭头）→ Word 原生三列时序表格
"""
import argparse
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FOOTER_LINE = 'TactileSense 研究文档 · 嵌入式设备组网配网流程'

doc = Document()
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)


def _songti(run, size=11, east='宋体'):
    run.font.name = east
    run.font.size = Pt(size)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), east)


def add_formatted_run(paragraph, text, size=11, italic=False, color=None, bold_all=False):
    tokens = re.split(r'(\*\*.*?\*\*|`[^`]+`)', text)
    for token in tokens:
        if token == '':
            continue
        if token.startswith('**') and token.endswith('**') and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`') and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(max(size - 1, 8))
            run.font.color.rgb = RGBColor(0xCC, 0x33, 0x33)
            run.bold = bold_all
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            if italic:
                run.italic = True
            if color:
                run.font.color.rgb = color
            continue
        else:
            run = paragraph.add_run(token)
            run.bold = bold_all
        _songti(run, size)
        run.italic = italic
        if color:
            run.font.color.rgb = color
    return paragraph


def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def set_col_width(table, col_idx, width_cm):
    for row in table.rows:
        row.cells[col_idx].width = Cm(width_cm)


def build_cover(title, version, date):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
    _songti(run, 26)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'{version} · {date}')
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    _songti(run, 14)
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(FOOTER_LINE)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    _songti(run, 11)
    doc.add_page_break()


def add_toc():
    h = doc.add_heading('目录', level=1)
    for run in h.runs:
        _songti(run, 16)
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = '目录将在 Word 中打开时自动更新（首次打开如提示更新域，请选"是"）。'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    for el in (fld_begin, instr, fld_sep, t, fld_end):
        run._r.append(el)
    doc.add_page_break()


# ---------------- ASCII 图 → Word 原生表格 ----------------

def render_grid_diagram(block_lines):
    """┌─┬─┐ 网格图（层次架构图/帧结构图）→ 带底纹的原生表格。"""
    cols = block_lines[0].count('┬') + 1
    rows = []
    cur = None
    for ln in block_lines[1:]:
        s = ln.strip()
        if s.startswith('└'):
            if cur is not None:
                rows.append(cur)
                cur = None
            break
        if s.startswith('├'):
            if cur is not None:
                rows.append(cur)
                cur = None
            continue
        if s.startswith('│'):
            parts = [p.strip() for p in s.split('│')[1:-1]]
            if cur is None:
                cur = parts
            else:
                for i, part in enumerate(parts):
                    if not part:
                        continue
                    if i < len(cur):
                        cur[i] = cur[i] + '\n' + part if cur[i] else part
                    else:
                        cur.append(part)
    if cur:
        rows.append(cur)
    if not rows:
        return False
    palette = ['D9E2F3', 'F2F2F2', 'E2EFDA', 'FFF2CC', 'FCE4D6', 'EDEDED']
    table = doc.add_table(rows=len(rows), cols=cols, style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_cells in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            set_cell_bg(cell, palette[r_idx % len(palette)])
            text = row_cells[c_idx] if c_idx < len(row_cells) else ''
            cell.text = ''
            first = True
            for seg in text.split('\n'):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                first = False
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                add_formatted_run(p, seg, size=10)
    if cols == 1:
        set_col_width(table, 0, 16.0)
    doc.add_paragraph()
    return True


CIRCLED = '①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳'


def render_sequence_diagram(block_lines):
    """│ 生命线 + →/← 箭头时序图 → 三列时序表格（host | 消息 | device）。"""
    first = block_lines[0].strip()
    titles = re.split(r'\s{2,}', first, maxsplit=1)
    t_left = titles[0].strip() if titles else 'host'
    t_right = titles[1].strip() if len(titles) > 1 else 'device'

    steps = []  # {'kind':'act'|'msg', 'left','right','label','arrow'}
    pending_left = None
    body = block_lines[1:]
    for idx, line in enumerate(body):
        if not line.strip():
            continue
        segs = line.split('│')
        left = segs[1].strip() if len(segs) > 1 else ''
        right = segs[2].strip() if len(segs) > 2 else ''
        has_arrow = '→' in line or '←' in line
        has_dash = re.search(r'─{3,}', line) is not None
        if has_arrow and has_dash:
            label = re.sub(r'[│─→←]', ' ', line)
            label = re.sub(r'\s+', ' ', label).strip()
            if pending_left:
                label = (pending_left + ' ' + label).strip() if label else pending_left
                pending_left = None
            steps.append({'kind': 'msg', 'arrow': '→' if '→' in line else '←',
                          'label': label, 'right': right})
        elif has_dash:
            # 无箭头的长横线行：上一条消息标签的续行
            cont = re.sub(r'[│─→←]', ' ', line)
            cont = re.sub(r'\s+', ' ', cont).strip()
            if steps and steps[-1]['kind'] == 'msg' and cont:
                steps[-1]['label'] = (steps[-1]['label'] + ' ' + cont).strip()
            elif cont:
                steps.append({'kind': 'act', 'left': cont, 'right': ''})
        else:
            # 动作标注行：前瞻判断是否为下一条裸箭头消息的标签
            is_msg_label = False
            if left and idx + 1 < len(body):
                nxt = body[idx + 1]
                nxt_label = re.sub(r'[│─→←]', ' ', nxt)
                nxt_label = re.sub(r'\s+', ' ', nxt_label).strip()
                if ('→' in nxt or '←' in nxt) and re.search(r'─{3,}', nxt) and not nxt_label:
                    is_msg_label = True
            if is_msg_label:
                pending_left = left
                if right:
                    steps.append({'kind': 'act', 'left': '', 'right': right})
            else:
                if left or right:
                    steps.append({'kind': 'act', 'left': left, 'right': right})

    table = doc.add_table(rows=len(steps) + 1, cols=3, style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    for c, text in enumerate((t_left, '消息', t_right)):
        cell = header.cells[c]
        set_cell_bg(cell, '1F4E79')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _songti(run, 10)
    for r_idx, step in enumerate(steps, start=1):
        row = table.rows[r_idx]
        if step['kind'] == 'msg':
            set_cell_bg(row.cells[0], 'F7F9FC')
            set_cell_bg(row.cells[2], 'F7F9FC')
            mid = row.cells[1]
            set_cell_bg(mid, 'FFF2CC')
            p = mid.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text = f'── {step["label"]} ──→' if step['arrow'] == '→' else f'←── {step["label"]} ──'
            add_formatted_run(p, text, size=10, bold_all=True)
            if step.get('right'):
                add_formatted_run(row.cells[2].paragraphs[0], step['right'], size=10)
        else:
            if step['left']:
                add_formatted_run(row.cells[0].paragraphs[0], step['left'], size=10)
            if step['right']:
                add_formatted_run(row.cells[2].paragraphs[0], step['right'], size=10)
    set_col_width(table, 0, 5.6)
    set_col_width(table, 1, 5.0)
    set_col_width(table, 2, 5.6)
    doc.add_paragraph()
    return True


def render_code_block(block_lines):
    for cl in block_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(1)
        run = p.add_run(cl if cl else ' ')
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


# ---------------- 主解析 ----------------

def convert(md_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    i = 0
    skip_h1 = True
    first_h2 = True
    while i < len(lines):
        line = lines[i].rstrip('\n')
        if line.strip() == '':
            i += 1
            continue
        if skip_h1 and re.match(r'^#\s+', line):
            skip_h1 = False
            i += 1
            continue

        # 代码块（含 ASCII 图识别）
        if line.strip().startswith('```'):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                block.append(lines[i].rstrip('\n'))
                i += 1
            i += 1
            stripped0 = block[0].strip() if block else ''
            joined = '\n'.join(block)
            if stripped0.startswith('┌') and '│' in joined:
                if not render_grid_diagram(block):
                    render_code_block(block)
            elif ('→' in joined or '←' in joined) and re.search(r'─{3,}', joined) and '│' in joined:
                if not render_sequence_diagram(block):
                    render_code_block(block)
            else:
                render_code_block(block)
            continue

        # 引用
        if line.lstrip().startswith('>'):
            content = re.sub(r'^>\s?', '', line.lstrip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(1)
            add_formatted_run(p, content, size=10, italic=True,
                              color=RGBColor(0x66, 0x66, 0x66))
            i += 1
            continue

        # Markdown 表格
        if (line.strip().startswith('|') and i + 1 < len(lines)
                and lines[i + 1].strip().startswith('|')
                and re.match(r'^[\|\s\-:]+$', lines[i + 1].strip())):
            header_line = line
            i += 2
            table_data = [[c.strip() for c in header_line.strip().strip('|').split('|')]]
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_data.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            num_cols = len(table_data[0])
            table = doc.add_table(rows=len(table_data), cols=num_cols,
                                  style='Light Grid Accent 1')
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row_data in enumerate(table_data):
                for c_idx in range(min(num_cols, len(row_data))):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = ''
                    add_formatted_run(cell.paragraphs[0], row_data[c_idx],
                                      size=10, bold_all=(r_idx == 0))
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.+)$', line)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            text = re.sub(r'`([^`]+)`', r'\1', text)
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            if level == 2:
                if not first_h2:
                    doc.add_page_break()
                first_h2 = False
            h = doc.add_heading(text, level=level)
            for run in h.runs:
                _songti(run, h.style.font.size.pt if h.style.font.size else 14)
            i += 1
            continue

        # 无序列表
        if re.match(r'^(\s*)[-*]\s+(.+)$', line):
            indent_count = len(re.match(r'^(\s*)', line).group(1))
            content = re.sub(r'^\s*[-*]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(1.5 * (indent_count // 2 + 1)) if indent_count else Cm(1)
            add_formatted_run(p, content)
            i += 1
            continue

        # 有序列表
        m_ol = re.match(r'^(\s*)(\d+)\.\s+(.+)$', line)
        if m_ol:
            indent_count = len(m_ol.group(1))
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Cm(1.5 * (indent_count // 2 + 1)) if indent_count else Cm(1)
            add_formatted_run(p, m_ol.group(3))
            i += 1
            continue

        # 分隔线
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run('—' * 40)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(10)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_formatted_run(p, line)
        i += 1


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('input_md')
    parser.add_argument('output_docx')
    parser.add_argument('--title', required=True)
    parser.add_argument('--version', required=True)
    parser.add_argument('--date', default='2026-08-05')
    args = parser.parse_args()

    build_cover(args.title, args.version, args.date)
    add_toc()
    convert(os.path.join(BASE_DIR, args.input_md) if not os.path.isabs(args.input_md)
            else args.input_md)
    out = args.output_docx if os.path.isabs(args.output_docx) \
        else os.path.join(BASE_DIR, args.output_docx)
    doc.save(out)
    print(f'Done: {out}')


if __name__ == '__main__':
    main()
