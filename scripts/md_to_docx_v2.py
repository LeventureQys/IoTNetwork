import re
import os
import sys
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_formatted_run(paragraph, text, size=11, italic=False, color=None, bold_all=False):
    tokens = re.split(r'(\*\*.*?\*\*|`[^`]+`)', text)
    for token in tokens:
        if token == '':
            continue
        if token.startswith('**') and token.endswith('**') and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = '宋体'
            run.font.size = Pt(size)
            run.italic = italic
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            if color:
                run.font.color.rgb = color
        elif token.startswith('`') and token.endswith('`') and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(max(size - 1, 8))
            run.font.color.rgb = RGBColor(0xCC, 0x33, 0x33)
            run.bold = bold_all
        else:
            run = paragraph.add_run(token)
            run.font.name = '宋体'
            run.font.size = Pt(size)
            run.italic = italic
            run.bold = bold_all
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            if color:
                run.font.color.rgb = color
    return paragraph


def setup_document(doc_title_str=None):
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    return doc


def add_cover(doc, title, subtitle=None):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.color.rgb = RGBColor(0x1A, 0x23, 0x7E)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(subtitle)
        run.font.size = Pt(13)
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('TactileSense · ESP32-C2 移植文档')
    run.font.size = Pt(10)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    doc.add_page_break()


def add_toc(doc):
    h = doc.add_heading('目录', level=1)
    for run in h.runs:
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = '（在 Word 中右键点击此处 → 更新域 → 更新整个目录）'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(t)
    run._r.append(fld_end)
    doc.add_page_break()


def convert_md_to_docx(md_path, docx_path, title=None, subtitle=None, base_dir=None):
    if base_dir is None:
        base_dir = os.path.dirname(md_path) if md_path else os.getcwd()

    doc = setup_document()

    if title:
        add_cover(doc, title, subtitle)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    skip_h1 = bool(title)   # if cover has title, skip md's first H1
    first_h2 = True
    has_toc = False

    while i < len(lines):
        line = lines[i].rstrip('\n')

        if line.strip() == '':
            i += 1
            continue

        # [TOC] placeholder
        if line.strip() == '[TOC]':
            add_toc(doc)
            has_toc = True
            i += 1
            continue

        # skip H1 if cover present
        if skip_h1 and re.match(r'^#\s+', line):
            skip_h1 = False
            i += 1
            continue

        # Image
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line.strip())
        if img_match:
            alt_text = img_match.group(1)
            rel_path = img_match.group(2)
            img_path = os.path.normpath(os.path.join(os.path.dirname(md_path), rel_path))
            if not os.path.exists(img_path):
                # try relative to base_dir
                img_path = os.path.normpath(os.path.join(base_dir, rel_path))
            if os.path.exists(img_path):
                try:
                    from PIL import Image
                    with Image.open(img_path) as im:
                        w_px, h_px = im.size
                    max_w_in, max_h_in = 5.5, 8.0
                    w_in = float(max_w_in)
                    h_in = w_in * h_px / w_px
                    if h_in > max_h_in:
                        h_in = float(max_h_in)
                        w_in = h_in * w_px / h_px
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(w_in), height=Inches(h_in))
                except Exception:
                    pass  # skip broken images silently
            i += 1
            # italic caption after image
            if i < len(lines) and re.match(r'^\*.*\*$', lines[i].strip()):
                caption = lines[i].strip().strip('*')
                cap_p = doc.add_paragraph()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_run = cap_p.add_run(caption)
                cap_run.font.size = Pt(9)
                cap_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                cap_run.italic = True
                cap_run.font.name = '宋体'
                cap_run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                i += 1
            doc.add_paragraph()
            continue

        # Code block
        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            i += 1
            for cl in code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(1)
                run = p.add_run(cl)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            continue

        # Blockquote
        if line.lstrip().startswith('> '):
            content = re.sub(r'^>\s?', '', line.lstrip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            m_list = re.match(r'^(\s*)[-*]\s+(.+)$', content)
            if m_list:
                p.style = 'List Bullet'
                indent_count = len(m_list.group(1))
                p.paragraph_format.left_indent = Cm(1 + 0.5 * (indent_count // 2 + 1))
                content = m_list.group(2)
            else:
                p.paragraph_format.left_indent = Cm(1)
            add_formatted_run(p, content, size=10, italic=True, color=RGBColor(0x66, 0x66, 0x66))
            i += 1
            continue

        # Table
        if i + 1 < len(lines) and lines[i + 1].strip().startswith('|') and re.match(r'^[\|\s\-:]+$', lines[i + 1].strip()):
            header_line = line
            i += 2
            table_data = []
            headers = [c.strip() for c in header_line.strip().strip('|').split('|')]
            table_data.append(headers)
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                table_data.append(row)
                i += 1
            num_cols = len(headers)
            table = doc.add_table(rows=len(table_data), cols=num_cols, style='Light Grid Accent 1')
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row_data in enumerate(table_data):
                for c_idx, cell_text in enumerate(row_data):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = ''
                    p = cell.paragraphs[0]
                    add_formatted_run(p, cell_text, size=9, bold_all=(r_idx == 0))
            doc.add_paragraph()
            continue

        # Heading
        match = re.match(r'^(#{1,4})\s+(.+)$', line)
        if match:
            level = len(match.group(1))
            text = match.group(2)
            if level == 2:
                if not first_h2:
                    doc.add_page_break()
                first_h2 = False
            h = doc.add_heading(text, level=level)
            for run in h.runs:
                run.font.name = '宋体'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            i += 1
            continue

        # Checkbox list: - [ ] or - [x]
        cb_match = re.match(r'^- \[([ x])\]\s+(.+)$', line)
        if cb_match:
            checked = cb_match.group(1) == 'x'
            checkbox_char = '☒ ' if checked else '☐ '
            text = checkbox_char + cb_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_formatted_run(p, text, size=10)
            i += 1
            continue

        # Unordered list
        if re.match(r'^(\s*)[-*]\s+(.+)$', line):
            indent_count = len(re.match(r'^(\s*)', line).group(1))
            content = re.sub(r'^\s*[-*]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            if indent_count > 0:
                p.paragraph_format.left_indent = Cm(1.5 * (indent_count // 2 + 1))
            else:
                p.paragraph_format.left_indent = Cm(1)
            add_formatted_run(p, content, size=10)
            i += 1
            continue

        # Ordered list
        match_ol = re.match(r'^(\s*)(\d+)\.\s+(.+)$', line)
        if match_ol:
            indent_count = len(match_ol.group(1))
            content = match_ol.group(3)
            p = doc.add_paragraph(style='List Number')
            if indent_count > 0:
                p.paragraph_format.left_indent = Cm(1.5 * (indent_count // 2 + 1))
            else:
                p.paragraph_format.left_indent = Cm(1)
            add_formatted_run(p, content, size=10)
            i += 1
            continue

        # Horizontal rule
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run('—' * 40)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            run.font.size = Pt(10)
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        add_formatted_run(p, line, size=10)
        i += 1

    doc.save(docx_path)
    return True


def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    md_dir = os.path.join(base_dir, 'Document', '移植说明', 'MD')
    word_dir = os.path.join(base_dir, 'Document', '移植说明', 'Word')

    tasks = [
        ('ESP32移植指南.md', 'ESP32移植指南.docx', 'ESP32 移植指南', '提供 ESP32-C2 后端移植的逐步骤任务清单'),
        ('ESP32移植方案.md', 'ESP32移植方案.docx', 'ESP32 移植方案', '硬件抽象层 29 个接口的实现方案与关键约定值'),
        ('配网流程说明.md', '配网流程说明.docx', '配网流程说明', '设备与上位机配网及连接保障流程'),
        ('验收说明.md', '验收说明.docx', '验收说明', 'ESP32-C2 移植完成后的自验清单'),
    ]

    for md_name, docx_name, title, subtitle in tasks:
        md_path = os.path.join(md_dir, md_name)
        docx_path = os.path.join(word_dir, docx_name)
        if not os.path.exists(md_path):
            print(f'SKIP: {md_path} not found')
            continue
        convert_md_to_docx(md_path, docx_path, title=title, subtitle=subtitle, base_dir=base_dir)
        print(f'Done: {docx_path}')


if __name__ == '__main__':
    main()
